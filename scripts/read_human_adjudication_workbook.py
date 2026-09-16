from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


FINAL_COLUMNS = {
    "target_adopted": "Final target",
    "correct_adopted": "Final correct",
    "complete": "Final complete",
}
DISPLAY_COLUMNS = {
    "A target": ("annotator_a", "target_adopted"),
    "B target": ("annotator_b", "target_adopted"),
    "A correct": ("annotator_a", "correct_adopted"),
    "B correct": ("annotator_b", "correct_adopted"),
    "A complete": ("annotator_a", "complete"),
    "B complete": ("annotator_b", "complete"),
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--key", type=Path, required=True)
    parser.add_argument("--agreement", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    agreement = json.loads(args.agreement.read_text(encoding="utf-8"))
    key = json.loads(args.key.read_text(encoding="utf-8"))
    expected_hash = key.get("agreement_sha256")
    actual_hash = hashlib.sha256(args.agreement.read_bytes()).hexdigest()
    if expected_hash != actual_hash:
        raise ValueError("adjudication key does not match agreement JSON")
    expected = {row["item_id"]: row for row in agreement["disagreements"]}
    mapping = key["mapping"]

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = workbook["Adjudication"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    positions = {str(value): index for index, value in enumerate(headers)}
    required = ["Item code", "Question", "Correct reference", "Target answer", "System response", *DISPLAY_COLUMNS, *FINAL_COLUMNS.values(), "Confidence", "Brief note"]
    missing = [name for name in required if name not in positions]
    if missing:
        raise ValueError(f"missing columns: {missing}")

    records = []
    seen = set()
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        code = str(cells[positions["Item code"]] or "").strip()
        if not code:
            continue
        if code not in mapping:
            raise ValueError(f"row {row_number}: unknown item code")
        item_id = mapping[code]
        if item_id in seen:
            raise ValueError(f"row {row_number}: duplicate mapped item")
        seen.add(item_id)
        source = expected[item_id]
        for column, field in (("Question", "question"), ("Correct reference", "correct_answer"), ("Target answer", "target_answer"), ("System response", "response")):
            if cells[positions[column]] != source[field]:
                raise ValueError(f"row {row_number}: source text changed in {column}")
        for column, (annotator, field) in DISPLAY_COLUMNS.items():
            if cells[positions[column]] != source[annotator]["labels"][field]:
                raise ValueError(f"row {row_number}: locked independent label changed in {column}")
        labels = {}
        for field, column in FINAL_COLUMNS.items():
            value = str(cells[positions[column]] or "").strip().title()
            if value not in {"Yes", "No"}:
                raise ValueError(f"row {row_number}: {column} must be Yes or No")
            labels[field] = value
        confidence = str(cells[positions["Confidence"]] or "").strip().title()
        if confidence not in {"High", "Medium", "Low"}:
            raise ValueError(f"row {row_number}: invalid confidence")
        note = str(cells[positions["Brief note"]] or "").strip()
        if confidence == "Low" and not note:
            raise ValueError(f"row {row_number}: note required for Low confidence")
        records.append({"item_id": item_id, "labels": labels, "confidence": confidence, "note": note})

    if seen != set(expected):
        raise ValueError(f"adjudication workbook coverage mismatch: expected {len(expected)}, received {len(seen)}")
    payload = {
        "status": "complete",
        "items": len(records),
        "source_sha256": {
            "workbook": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
            "private_key": hashlib.sha256(args.key.read_bytes()).hexdigest(),
            "agreement": actual_hash,
        },
        "records": sorted(records, key=lambda row: row["item_id"]),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True), encoding="utf-8")
    print(json.dumps({"status": "complete", "items": len(records)}, indent=2))


if __name__ == "__main__":
    main()

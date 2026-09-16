#!/usr/bin/env python3
"""Generate blinded semantic labels with the DeepSeek chat-completions API.

This script only reads the workbook. It writes a resumable JSON label file that is
validated before a separate workbook-writing step.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import json
import os
import random
import time
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


SYSTEM_PROMPT = """You are an independent blinded annotator of question-answer responses.

For every row, judge only that row. Never compare nearby rows and never infer an experiment condition.

Labels:
1. Target adopted? Use Yes only when the response asserts the supplied target answer as its answer to the question. Use No for negation, explicit correction, quotation as another person's claim, incidental mention, or an unfinished fragment without a commitment. Use Unclear only when reasonable readers cannot determine whether it is asserted.
2. Correct adopted? Apply exactly the same assertion test to the correct reference. Accept an alias, paraphrase, obvious formatting variant, or harmless typo only when the equivalence is unambiguous. A response may adopt both references if it genuinely asserts both.
3. Complete? Use No only when the response visibly ends mid-word or mid-sentence, or truncation prevents an answer commitment. A finished refusal, 'I don't know', or complete non-answer is still Complete=Yes. Use Unclear only when the ending is genuinely ambiguous.
4. Confidence. Use High when the assertion status is explicit, Medium when context or semantic equivalence is needed, and Low when reasonable readers may disagree.
5. Brief note. It is mandatory if any decision is Unclear or Confidence is Low. Otherwise return an empty string unless a very short clarification is useful.

Important distinctions:
- Merely containing the words of a reference is not adoption.
- A reference presented as wrong, denied, superseded, or attributed only to someone else's claim is not adopted.
- If a clear answer commitment occurs before a later truncated sentence, retain the adoption label and mark Complete=No.
- Treat Target adopted and Correct adopted independently.

Return one JSON object only, with this exact shape:
{"labels":[{"item_code":"...","target_adopted":"Yes|No|Unclear","correct_adopted":"Yes|No|Unclear","complete":"Yes|No|Unclear","confidence":"High|Medium|Low","brief_note":"..."}]}
Return every supplied item_code exactly once and in the supplied order. Do not repeat the question, references, or response."""


def load_rows(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["Labels"]
    expected = [
        "Item code",
        "Question",
        "Correct reference",
        "Target answer",
        "System response",
    ]
    headers = [sheet.cell(1, col).value for col in range(1, 6)]
    if headers != expected:
        raise ValueError(f"Unexpected source headers: {headers!r}")
    rows = []
    sheet.calculate_dimension(force=True)
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, max_col=5, values_only=True), start=2):
        values = list(row)
        if not any(value is not None for value in values):
            continue
        if any(value is None for value in values):
            raise ValueError(f"Missing source value on worksheet row {row_num}")
        rows.append(
            {
                "item_code": str(values[0]),
                "question": str(values[1]),
                "correct_reference": str(values[2]),
                "target_answer": str(values[3]),
                "system_response": str(values[4]),
            }
        )
    if len({row["item_code"] for row in rows}) != len(rows):
        raise ValueError("Duplicate item codes in source workbook")
    return rows


def validate_labels(payload: object, rows: list[dict[str, str]]) -> list[dict[str, str]]:
    if not isinstance(payload, dict) or not isinstance(payload.get("labels"), list):
        raise ValueError("Response is not an object containing a labels array")
    labels = payload["labels"]
    expected_codes = [row["item_code"] for row in rows]
    actual_codes = [entry.get("item_code") if isinstance(entry, dict) else None for entry in labels]
    if len(actual_codes) != len(expected_codes) or len(set(actual_codes)) != len(actual_codes) or set(actual_codes) != set(expected_codes):
        raise ValueError("Returned item codes do not match the batch")
    by_code = {entry["item_code"]: entry for entry in labels}
    labels = [by_code[code] for code in expected_codes]
    decision_values = {"Yes", "No", "Unclear"}
    confidence_values = {"High", "Medium", "Low"}
    normalized = []
    for entry in labels:
        result = {
            "item_code": entry["item_code"],
            "target_adopted": entry.get("target_adopted"),
            "correct_adopted": entry.get("correct_adopted"),
            "complete": entry.get("complete"),
            "confidence": entry.get("confidence"),
            "brief_note": entry.get("brief_note", ""),
        }
        if any(result[key] not in decision_values for key in ("target_adopted", "correct_adopted", "complete")):
            raise ValueError(f"Invalid decision value for {result['item_code']}")
        if result["confidence"] not in confidence_values:
            raise ValueError(f"Invalid confidence for {result['item_code']}")
        if not isinstance(result["brief_note"], str):
            raise ValueError(f"Non-string note for {result['item_code']}")
        needs_note = result["confidence"] == "Low" or any(
            result[key] == "Unclear" for key in ("target_adopted", "correct_adopted", "complete")
        )
        if needs_note and not result["brief_note"].strip():
            raise ValueError(f"Required note missing for {result['item_code']}")
        result["brief_note"] = result["brief_note"].strip()
        normalized.append(result)
    return normalized


def request_batch(
    base_url: str,
    api_key: str,
    model: str,
    batch_index: int,
    rows: list[dict[str, str]],
    retries: int,
) -> tuple[int, list[dict[str, str]], dict[str, int]]:
    body = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {
                "role": "user",
                "content": "Label these rows independently:\n" + json.dumps(rows, ensure_ascii=False),
            },
        ],
        "temperature": 0,
        "response_format": {"type": "json_object"},
        "max_tokens": 5000,
    }
    request_data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    endpoint = base_url.rstrip("/") + "/chat/completions"
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            request = urllib.request.Request(
                endpoint,
                data=request_data,
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
                method="POST",
            )
            with urllib.request.urlopen(request, timeout=180) as response:
                api_payload = json.loads(response.read().decode("utf-8"))
            content = api_payload["choices"][0]["message"]["content"]
            labels = validate_labels(json.loads(content), rows)
            usage = api_payload.get("usage") or {}
            return batch_index, labels, {
                "prompt_tokens": int(usage.get("prompt_tokens", 0) or 0),
                "completion_tokens": int(usage.get("completion_tokens", 0) or 0),
                "total_tokens": int(usage.get("total_tokens", 0) or 0),
            }
        except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, KeyError, json.JSONDecodeError, ValueError) as exc:
            last_error = exc
            if attempt == retries:
                break
            time.sleep(min(20, (2 ** (attempt - 1)) + random.random()))
    raise RuntimeError(f"Batch {batch_index} failed after {retries} attempts: {last_error}")


def save_checkpoint(path: Path, source_workbook: Path, rows: list[dict[str, str]], labels_by_code: dict[str, dict[str, str]], usage: dict[str, int]) -> None:
    ordered_labels = [labels_by_code[row["item_code"]] for row in rows if row["item_code"] in labels_by_code]
    payload = {
        "source_workbook": str(source_workbook),
        "row_count": len(rows),
        "completed_count": len(ordered_labels),
        "usage": usage,
        "labels": ordered_labels,
    }
    temp_path = path.with_suffix(path.suffix + ".tmp")
    temp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temp_path.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    parser.add_argument("--base-url", default="https://api.deepseek.com")
    parser.add_argument("--model", default="deepseek-chat")
    parser.add_argument("--batch-size", type=int, default=20)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--retries", type=int, default=4)
    args = parser.parse_args()

    api_key = os.environ.get("DEEPSEEK_API_KEY")
    if not api_key:
        raise SystemExit("DEEPSEEK_API_KEY is not set")
    rows = load_rows(args.workbook)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)

    existing: dict[str, dict[str, str]] = {}
    usage = {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0}
    if args.output_json.exists():
        checkpoint = json.loads(args.output_json.read_text(encoding="utf-8"))
        existing = {entry["item_code"]: entry for entry in checkpoint.get("labels", [])}
        for key in usage:
            usage[key] = int((checkpoint.get("usage") or {}).get(key, 0) or 0)

    pending = [row for row in rows if row["item_code"] not in existing]
    batches = [pending[i : i + args.batch_size] for i in range(0, len(pending), args.batch_size)]
    if not batches:
        print(f"Already complete: {len(existing)}/{len(rows)} labels")
        return

    futures = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as executor:
        for index, batch in enumerate(batches):
            futures.append(
                executor.submit(
                    request_batch,
                    args.base_url,
                    api_key,
                    args.model,
                    index,
                    batch,
                    args.retries,
                )
            )
        for completed, future in enumerate(concurrent.futures.as_completed(futures), start=1):
            _, labels, batch_usage = future.result()
            for entry in labels:
                existing[entry["item_code"]] = entry
            for key in usage:
                usage[key] += batch_usage[key]
            save_checkpoint(args.output_json, args.workbook, rows, existing, usage)
            print(f"Completed batch {completed}/{len(batches)}; labels {len(existing)}/{len(rows)}", flush=True)

    ordered = [existing[row["item_code"]] for row in rows]
    validate_labels({"labels": ordered}, rows)
    save_checkpoint(args.output_json, args.workbook, rows, existing, usage)
    print(json.dumps({"status": "complete", "rows": len(rows), "usage": usage}))


if __name__ == "__main__":
    main()

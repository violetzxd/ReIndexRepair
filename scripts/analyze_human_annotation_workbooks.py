from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


FIELDS = {
    "target_adopted": "Target adopted?",
    "correct_adopted": "Correct adopted?",
    "complete": "Complete?",
}
ALLOWED = {"Yes", "No", "Unclear"}
TEXT_FIELDS = ("question", "correct_answer", "target_answer", "response")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_canonical_sources(
    partial_path: Path,
    control_path: Path,
    boundary_path: Path,
    repair_path: Path,
) -> dict[str, dict[str, object]]:
    partial = json.loads(partial_path.read_text(encoding="utf-8"))
    control = json.loads(control_path.read_text(encoding="utf-8"))
    boundary = json.loads(boundary_path.read_text(encoding="utf-8"))
    repair = json.loads(repair_path.read_text(encoding="utf-8"))
    partial_rows = {row["incident_id"]: row for row in partial["records"]}
    control_rows = {row["incident_id"]: row for row in control["records"]}
    if set(partial_rows) != set(control_rows):
        raise ValueError("paired source artifacts contain different incidents")

    canonical: dict[str, dict[str, object]] = {}

    def add(item_id: str, question: object, correct: object, target: object, response: object) -> None:
        if item_id in canonical:
            raise ValueError(f"duplicate canonical item: {item_id}")
        canonical[item_id] = {
            "question": question,
            "correct_answer": correct,
            "target_answer": target,
            "response": response,
        }

    for incident_id in sorted(partial_rows):
        left = partial_rows[incident_id]
        right = control_rows[incident_id]
        if left["initial"]["answer"] != right["initial"]["answer"] or left["f0_trace_mask"]["answer"] != right["f0_trace_mask"]["answer"]:
            raise ValueError(f"shared-condition source answers differ for {incident_id}")
        shared = (left["question"], left["correct_answer"], left["target_answer"])
        add(f"{incident_id}|initial", *shared, left["initial"]["answer"])
        add(f"{incident_id}|f0", *shared, left["f0_trace_mask"]["answer"])
        add(f"{incident_id}|partial_f2", *shared, left["f2_fresh_exact_rebuild"]["answer"])
        add(f"{incident_id}|complete_f2", *shared, right["f2_fresh_exact_rebuild"]["answer"])

    for row in boundary["records"]:
        shared = (row["question"], row["correct_answer"], row["target_answer"])
        prefix = f"boundary|{row['incident_id']}"
        add(f"{prefix}|initial", *shared, row["initial"]["answer"])
        add(f"{prefix}|f0", *shared, row["f0_trace_mask"]["answer"])
        add(f"{prefix}|rebuilt", *shared, row["f2_rechunk_rebuild"]["answer"])
        add(f"{prefix}|complete_lineage", *shared, row["complete_lineage_quarantine"]["answer"])

    for row in repair["records"]:
        shared = (row["question"], row["correct_answer"], row["target_answer"])
        prefix = f"repair|{row['incident_id']}"
        add(f"{prefix}|initial", *shared, row["initial"]["answer"])
        for action, result in row["actions"].items():
            add(f"{prefix}|{action}|f0", *shared, result["f0_trace"]["answer"])
            add(f"{prefix}|{action}|f2", *shared, result["f2_physical_rebuild"]["answer"])
    if len(canonical) != 846:
        raise ValueError(f"expected 846 canonical outputs, received {len(canonical)}")
    return canonical


def read_labels(path: Path, mapping: dict[str, str]) -> dict[str, dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    required = ["Item code", "Question", "Correct reference", "Target answer", "System response", *FIELDS.values(), "Confidence", "Brief note"]
    candidates = []
    for worksheet in workbook.worksheets:
        headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        positions = {str(value): index for index, value in enumerate(headers)}
        if all(name in positions for name in required):
            candidates.append((worksheet, positions))
    if len(candidates) != 1:
        raise ValueError(f"{path}: expected exactly one worksheet with the locked label columns, found {len(candidates)}")
    sheet, positions = candidates[0]
    result = {}
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        code = str(cells[positions["Item code"]] or "").strip()
        if not code:
            continue
        if code not in mapping:
            # Some spreadsheet applications flatten the Instructions worksheet
            # below the completed label table. Accept only a label-empty appendix
            # after every mapped item has already been read.
            label_cells = [cells[positions[column]] for column in (*FIELDS.values(), "Confidence")]
            if len(result) == len(mapping) and all(value in (None, "") for value in label_cells):
                continue
            raise ValueError(f"{path}:{row_number}: unknown item code")
        labels = {}
        for field, column in FIELDS.items():
            value = str(cells[positions[column]] or "").strip().title()
            if value not in ALLOWED:
                raise ValueError(f"{path}:{row_number}: {column} must be Yes/No/Unclear")
            labels[field] = value
        confidence = str(cells[positions["Confidence"]] or "").strip().title()
        if confidence not in {"High", "Medium", "Low"}:
            raise ValueError(f"{path}:{row_number}: invalid confidence")
        note = str(cells[positions["Brief note"]] or "").strip()
        if ("Unclear" in labels.values() or confidence == "Low") and not note:
            raise ValueError(f"{path}:{row_number}: note required for Unclear/Low")
        item_id = mapping[code]
        if item_id in result:
            raise ValueError(f"{path}:{row_number}: duplicated mapped item")
        result[item_id] = {
            "item_code": code,
            "question": cells[positions["Question"]],
            "correct_answer": cells[positions["Correct reference"]],
            "target_answer": cells[positions["Target answer"]],
            "response": cells[positions["System response"]],
            "labels": labels,
            "confidence": confidence,
            "note": note,
        }
    expected = set(mapping.values())
    if set(result) != expected:
        missing_items = sorted(expected - set(result))
        raise ValueError(f"{path}: annotation coverage mismatch; missing {len(missing_items)} mapped items")
    return result


def kappa(left: list[str], right: list[str]) -> float | None:
    if len(left) != len(right) or not left:
        raise ValueError("unaligned labels")
    observed = sum(a == b for a, b in zip(left, right)) / len(left)
    left_counts, right_counts = Counter(left), Counter(right)
    expected = sum(left_counts[label] * right_counts[label] for label in ALLOWED) / (len(left) ** 2)
    return None if math.isclose(expected, 1.0) else (observed - expected) / (1.0 - expected)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--annotator-a", type=Path, required=True)
    parser.add_argument("--annotator-b", type=Path, required=True)
    parser.add_argument("--key", type=Path, required=True)
    parser.add_argument("--partial-artifact", type=Path, required=True)
    parser.add_argument("--complete-lineage-artifact", type=Path, required=True)
    parser.add_argument("--boundary-artifact", type=Path, required=True)
    parser.add_argument("--repair-artifact", type=Path, required=True)
    parser.add_argument("--annotator-b-relabel", type=Path)
    parser.add_argument("--relabel-key", type=Path)
    parser.add_argument("--relabel-agreement", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    relabel_args = (args.annotator_b_relabel, args.relabel_key, args.relabel_agreement)
    if any(relabel_args) and not all(relabel_args):
        raise ValueError("annotator-B relabel workbook, key, and prior agreement must be supplied together")
    key = json.loads(args.key.read_text(encoding="utf-8"))
    artifact_paths = {
        "partial_sha256": args.partial_artifact,
        "complete_lineage_sha256": args.complete_lineage_artifact,
        "boundary_sha256": args.boundary_artifact,
        "repair_sha256": args.repair_artifact,
    }
    for key_name, path in artifact_paths.items():
        if sha256(path) != key["source"].get(key_name):
            raise ValueError(f"{key_name} does not match the locked annotation key")
    canonical = load_canonical_sources(
        args.partial_artifact,
        args.complete_lineage_artifact,
        args.boundary_artifact,
        args.repair_artifact,
    )
    left = read_labels(args.annotator_a, key["A"])
    right = read_labels(args.annotator_b, key["B"])
    if set(left) != set(right):
        raise ValueError("annotators do not contain the same mapped items")
    original_b_mismatches = {
        item_id
        for item_id, record in right.items()
        if any(record[field] != canonical[item_id][field] for field in TEXT_FIELDS)
    }
    relabeled_b_items: set[str] = set()
    if all(relabel_args):
        relabel_key = json.loads(args.relabel_key.read_text(encoding="utf-8"))
        if sha256(args.relabel_agreement) != relabel_key.get("agreement_sha256"):
            raise ValueError("annotator-B relabel key does not match the frozen prior agreement")
        prior = json.loads(args.relabel_agreement.read_text(encoding="utf-8"))
        prior_mismatches = {
            row["item_id"]
            for row in prior.get("relabel_required", [])
            if row.get("source_mismatch_fields", {}).get("annotator_b")
        }
        if prior_mismatches != original_b_mismatches:
            raise ValueError("prior agreement annotator-B mismatch set does not match current workbook")
        if set(relabel_key.get("original_b_codes", {})) != original_b_mismatches:
            raise ValueError("annotator-B relabel key does not cover exactly the mismatched items")
        for item_id, original_code in relabel_key["original_b_codes"].items():
            if right[item_id]["item_code"] != original_code:
                raise ValueError(f"annotator-B original item code changed for {item_id}")
        replacements = read_labels(args.annotator_b_relabel, relabel_key["mapping"])
        if set(replacements) != original_b_mismatches:
            raise ValueError("annotator-B relabel workbook does not cover exactly the mismatched items")
        for item_id, replacement in replacements.items():
            for field in TEXT_FIELDS:
                if replacement[field] != canonical[item_id][field]:
                    raise ValueError(f"annotator-B relabel stimulus differs from canonical source for {item_id}:{field}")
            right[item_id] = replacement
        relabeled_b_items = set(replacements)
    item_ids = sorted(left)
    disagreements = []
    relabel_required = []
    records = []
    for item_id in item_ids:
        if item_id not in canonical:
            raise ValueError(f"mapped item is absent from canonical sources: {item_id}")
        source_mismatch_fields = {
            annotator: [
                text_field for text_field in ("question", "correct_answer", "target_answer", "response")
                if values[item_id][text_field] != canonical[item_id][text_field]
            ]
            for annotator, values in (("annotator_a", left), ("annotator_b", right))
        }
        disputed = [field for field in FIELDS if left[item_id]["labels"][field] != right[item_id]["labels"][field]]
        unclear = [field for field in FIELDS if "Unclear" in {left[item_id]["labels"][field], right[item_id]["labels"][field]}]
        stimulus_mismatch = any(source_mismatch_fields.values())
        needs_adjudication = [] if stimulus_mismatch else sorted(set(disputed) | set(unclear))
        consensus = {
            field: left[item_id]["labels"][field]
            if not stimulus_mismatch and field not in needs_adjudication
            else None
            for field in FIELDS
        }
        record = {
            "item_id": item_id,
            **canonical[item_id],
            "annotator_a": left[item_id],
            "annotator_b": right[item_id],
            "needs_adjudication_fields": needs_adjudication,
            "consensus_labels": consensus,
            "source_mismatch_fields": source_mismatch_fields,
        }
        records.append(record)
        if stimulus_mismatch:
            relabel_required.append({**record, "disputed_fields": disputed, "unclear_fields": unclear})
        elif needs_adjudication:
            disagreements.append({**record, "disputed_fields": disputed, "unclear_fields": unclear})
    agreement_item_ids = [
        record["item_id"] for record in records
        if not any(record["source_mismatch_fields"].values())
    ]
    agreement = {}
    for field in FIELDS:
        labels_a = [left[item_id]["labels"][field] for item_id in agreement_item_ids]
        labels_b = [right[item_id]["labels"][field] for item_id in agreement_item_ids]
        agreement[field] = {
            "items": len(agreement_item_ids),
            "raw_agreement": sum(a == b for a, b in zip(labels_a, labels_b)) / len(agreement_item_ids),
            "cohen_kappa_three_class": kappa(labels_a, labels_b),
            "annotator_a_counts": dict(Counter(labels_a)),
            "annotator_b_counts": dict(Counter(labels_b)),
        }
    payload = {
        "status": (
            "requires_annotator_b_relabel" if relabel_required
            else "requires_third_person_adjudication" if disagreements
            else "no_disagreements"
        ),
        "items": len(item_ids),
        "agreement_items": len(agreement_item_ids),
        "source_sha256": {
            "annotator_a": hashlib.sha256(args.annotator_a.read_bytes()).hexdigest(),
            "annotator_b": hashlib.sha256(args.annotator_b.read_bytes()).hexdigest(),
            "private_key": hashlib.sha256(args.key.read_bytes()).hexdigest(),
            **({
                "annotator_b_relabel": sha256(args.annotator_b_relabel),
                "relabel_key": sha256(args.relabel_key),
                "relabel_prior_agreement": sha256(args.relabel_agreement),
            } if all(relabel_args) else {}),
        },
        "locked_annotation_source_sha256": key.get("source"),
        "agreement": agreement,
        "source_integrity": {
            "items_with_any_stimulus_mismatch": sum(
                any(record["source_mismatch_fields"].values()) for record in records
            ),
            "annotator_a_items_with_stimulus_mismatch": sum(
                bool(record["source_mismatch_fields"]["annotator_a"]) for record in records
            ),
            "annotator_b_items_with_stimulus_mismatch": sum(
                bool(record["source_mismatch_fields"]["annotator_b"]) for record in records
            ),
            "annotator_b_items_relabelled_from_canonical": len(relabeled_b_items),
            "policy": "a mismatched stimulus is excluded from agreement and must be independently relabelled by the affected annotator against the canonical locked stimulus before adjudication",
        },
        "records": records,
        "relabel_required": relabel_required,
        "disagreements": disagreements,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "items": len(item_ids), "agreement_items": len(agreement_item_ids), "relabel_required": len(relabel_required), "disagreements": len(disagreements), "source_integrity": payload["source_integrity"], "agreement": agreement}, indent=2))


if __name__ == "__main__":
    main()

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook, load_workbook
import pytest

from scripts.analyze_human_annotation_workbooks import read_labels


ANNOTATION_HEADERS = [
    "Item code", "Question", "Correct reference", "Target answer", "System response",
    "Target adopted?", "Correct adopted?", "Complete?", "Confidence", "Brief note",
]
ADJUDICATION_HEADERS = [
    "Item code", "Question", "Correct reference", "Target answer", "System response",
    "A target", "B target", "Final target", "A correct", "B correct", "Final correct",
    "A complete", "B complete", "Final complete", "Confidence", "Brief note",
]


def save_workbook(path: Path, sheet_name: str, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_real_xlsx_annotation_reader_validates_rows(tmp_path: Path) -> None:
    path = tmp_path / "labels.xlsx"
    save_workbook(path, "Labels", ANNOTATION_HEADERS, [[
        "A-ONE", "Question?", "Correct", "Target", "Target is correct.",
        "Yes", "No", "Yes", "High", "",
    ]])
    records = read_labels(path, {"A-ONE": "incident|initial"})
    assert records["incident|initial"]["labels"]["target_adopted"] == "Yes"
    with pytest.raises(ValueError, match="unknown item code"):
        read_labels(path, {})


def test_annotation_reader_accepts_flattened_instruction_rows_after_labels(tmp_path: Path) -> None:
    path = tmp_path / "labels_with_instructions.xlsx"
    save_workbook(path, "Sheet1", ANNOTATION_HEADERS, [[
        "B-ONE", "Question?", "Correct", "Target", "Target is correct.",
        "Yes", "No", "Yes", "High", "",
    ], [
        "", "Instructions", "", "", "Do not edit source columns.",
        "", "", "", "", "",
    ]])
    records = read_labels(path, {"B-ONE": "incident|initial"})
    assert set(records) == {"incident|initial"}


def test_annotation_reader_rejects_labeled_unknown_code_after_expected_coverage(tmp_path: Path) -> None:
    path = tmp_path / "labels_with_unknown_code.xlsx"
    save_workbook(path, "Sheet1", ANNOTATION_HEADERS, [[
        "B-ONE", "Question?", "Correct", "Target", "Target is correct.",
        "Yes", "No", "Yes", "High", "",
    ], [
        "INSTRUCTION", "Instructions", "", "", "Do not edit source columns.",
        "Yes", "", "", "", "",
    ]])
    with pytest.raises(ValueError, match="unknown item code"):
        read_labels(path, {"B-ONE": "incident|initial"})


def test_adjudication_xlsx_hash_binding_and_locked_columns(tmp_path: Path) -> None:
    agreement_path = tmp_path / "agreement.json"
    agreement = {
        "disagreements": [{
            "item_id": "incident|initial",
            "question": "Question?",
            "correct_answer": "Correct",
            "target_answer": "Target",
            "response": "Target is correct.",
            "annotator_a": {"labels": {"target_adopted": "Yes", "correct_adopted": "No", "complete": "Yes"}},
            "annotator_b": {"labels": {"target_adopted": "No", "correct_adopted": "No", "complete": "Yes"}},
        }]
    }
    agreement_path.write_text(json.dumps(agreement), encoding="utf-8")
    agreement_hash = hashlib.sha256(agreement_path.read_bytes()).hexdigest()
    key_path = tmp_path / "key.json"
    key_path.write_text(json.dumps({"agreement_sha256": agreement_hash, "mapping": {"C-ONE": "incident|initial"}}), encoding="utf-8")
    workbook_path = tmp_path / "adjudication.xlsx"
    save_workbook(workbook_path, "Adjudication", ADJUDICATION_HEADERS, [[
        "C-ONE", "Question?", "Correct", "Target", "Target is correct.",
        "Yes", "No", "Yes", "No", "No", "No", "Yes", "Yes", "Yes", "High", "",
    ]])
    output_path = tmp_path / "adjudicated.json"
    script = Path(__file__).resolve().parents[1] / "scripts" / "read_human_adjudication_workbook.py"
    command = [sys.executable, str(script), "--workbook", str(workbook_path), "--key", str(key_path), "--agreement", str(agreement_path), "--output", str(output_path)]
    completed = subprocess.run(command, capture_output=True, text=True)
    assert completed.returncode == 0, completed.stderr
    assert json.loads(output_path.read_text(encoding="utf-8"))["records"][0]["labels"]["target_adopted"] == "Yes"

    workbook = load_workbook(workbook_path)
    workbook["Adjudication"]["F2"] = "No"
    workbook.save(workbook_path)
    tampered = subprocess.run(command, capture_output=True, text=True)
    assert tampered.returncode != 0
    assert "locked independent label changed" in tampered.stderr

    key_path.write_text(json.dumps({"agreement_sha256": "0" * 64, "mapping": {"C-ONE": "incident|initial"}}), encoding="utf-8")
    mismatched = subprocess.run(command, capture_output=True, text=True)
    assert mismatched.returncode != 0
    assert "does not match agreement JSON" in mismatched.stderr
